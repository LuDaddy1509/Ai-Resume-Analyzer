from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from io import BytesIO
from datetime import datetime
from typing import Dict, Any, List
from sqlalchemy.orm import Session


class ExportService:
    @staticmethod
    def generate_pdf(analysis_data: Dict[str, Any]) -> BytesIO:
        """Generate PDF report from analysis data."""
        buffer = BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=72,
            leftMargin=72,
            topMargin=72,
            bottomMargin=36
        )

        story = []
        styles = getSampleStyleSheet()

        # Custom styles
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=24,
            textColor=colors.HexColor('#2c3e50'),
            spaceAfter=30,
            alignment=1  # Center
        )

        heading_style = ParagraphStyle(
            'CustomHeading',
            parent=styles['Heading2'],
            fontSize=16,
            textColor=colors.HexColor('#3498db'),
            spaceAfter=12,
            spaceBefore=20
        )

        # Header
        story.append(Paragraph("AI Resume Analyzer", title_style))
        story.append(Paragraph("Analysis Report", styles['Heading2']))
        story.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles['Normal']))
        story.append(Spacer(1, 0.3 * inch))

        # CV Information
        story.append(Paragraph("Resume Information", heading_style))
        cv_data = [
            ['File Name:', analysis_data.get('filename', 'N/A')],
            ['Candidate:', analysis_data.get('full_name', 'N/A')],
            ['Email:', analysis_data.get('email', 'N/A')],
            ['Phone:', analysis_data.get('phone', 'N/A') or 'N/A'],
            ['Analysis Date:', analysis_data.get('created_at', 'N/A')]
        ]
        cv_table = Table(cv_data, colWidths=[2 * inch, 4 * inch])
        cv_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.lightgrey),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 1, colors.grey)
        ]))
        story.append(cv_table)
        story.append(Spacer(1, 0.3 * inch))

        # Job Description Info (if available)
        if analysis_data.get('job_description_snapshot'):
            story.append(Paragraph("Job Description", heading_style))
            jd_info = analysis_data.get('job_description_snapshot', {})
            jd_data = [
                ['Position:', jd_info.get('title', 'N/A')],
                ['Company:', jd_info.get('company', 'N/A') or 'N/A'],
                ['Level:', jd_info.get('job_level', 'N/A') or 'N/A']
            ]
            jd_table = Table(jd_data, colWidths=[2 * inch, 4 * inch])
            jd_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (0, -1), colors.lightgrey),
                ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 10),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
                ('GRID', (0, 0), (-1, -1), 1, colors.grey)
            ]))
            story.append(jd_table)
            story.append(Spacer(1, 0.3 * inch))

        # Scores
        story.append(Paragraph("Analysis Scores", heading_style))
        match_score = analysis_data.get('match_score', 0)
        score_data = [
            ['Match Score:', f"{match_score}%"]
        ]
        score_table = Table(score_data, colWidths=[2 * inch, 4 * inch])
        score_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.lightgrey),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 1, colors.grey)
        ]))
        story.append(score_table)
        story.append(Spacer(1, 0.3 * inch))

        # Skills
        skills = analysis_data.get('skills', [])
        if skills:
            story.append(Paragraph("Skills Found", heading_style))
            skills_text = ', '.join(skills)
            story.append(Paragraph(skills_text, styles['Normal']))
            story.append(Spacer(1, 0.2 * inch))

        # Matched/Missing Keywords (from analysis)
        analysis = analysis_data.get('analysis', {})
        if analysis:
            matched = analysis.get('matched_skills', [])
            if matched:
                story.append(Paragraph("Matched Keywords", heading_style))
                matched_text = ', '.join(matched)
                story.append(Paragraph(matched_text, styles['Normal']))
                story.append(Spacer(1, 0.2 * inch))

            missing = analysis.get('missing_skills', [])
            if missing:
                story.append(Paragraph("Missing Keywords", heading_style))
                missing_text = ', '.join(missing)
                story.append(Paragraph(missing_text, styles['Normal']))
                story.append(Spacer(1, 0.2 * inch))

            # Suggestions
            suggestions = analysis.get('suggestions', [])
            if suggestions:
                story.append(Paragraph("Recommendations", heading_style))
                for i, suggestion in enumerate(suggestions[:5], 1):
                    story.append(Paragraph(f"{i}. {suggestion}", styles['Normal']))
                story.append(Spacer(1, 0.2 * inch))

        # Footer disclaimer
        story.append(Spacer(1, 0.5 * inch))
        disclaimer_style = ParagraphStyle(
            'Disclaimer',
            parent=styles['Normal'],
            fontSize=8,
            textColor=colors.grey,
            alignment=1
        )
        story.append(Paragraph(
            "Note: This analysis is for reference purposes only. "
            "Users should verify all information before use. "
            "The system does not guarantee recruitment outcomes.",
            disclaimer_style
        ))

        # Build PDF
        doc.build(story)
        buffer.seek(0)
        return buffer

    @staticmethod
    def generate_excel_single(analysis_data: Dict[str, Any]) -> BytesIO:
        """Generate Excel report for a single resume analysis."""
        wb = Workbook()

        # Remove default sheet
        if 'Sheet' in wb.sheetnames:
            wb.remove(wb['Sheet'])

        # Sheet 1: Overview
        ws_overview = wb.create_sheet('Overview')

        # Header style
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')

        ws_overview['A1'] = 'Field'
        ws_overview['B1'] = 'Value'
        ws_overview['A1'].fill = header_fill
        ws_overview['A1'].font = header_font
        ws_overview['B1'].fill = header_fill
        ws_overview['B1'].font = header_font

        # Data
        overview_data = [
            ['Resume Name', analysis_data.get('filename', 'N/A')],
            ['Candidate Name', analysis_data.get('full_name', 'N/A')],
            ['Email', analysis_data.get('email', 'N/A')],
            ['Phone', analysis_data.get('phone', 'N/A') or 'N/A'],
            ['Analysis Date', analysis_data.get('created_at', 'N/A')],
            ['Match Score', f"{analysis_data.get('match_score', 0)}%"]
        ]

        for idx, row in enumerate(overview_data, start=2):
            ws_overview[f'A{idx}'] = row[0]
            ws_overview[f'B{idx}'] = row[1]
            ws_overview[f'A{idx}'].font = Font(bold=True)

        ws_overview.column_dimensions['A'].width = 20
        ws_overview.column_dimensions['B'].width = 40

        # Sheet 2: Skills
        ws_skills = wb.create_sheet('Skills')
        ws_skills['A1'] = 'Skill'
        ws_skills['A1'].fill = header_fill
        ws_skills['A1'].font = header_font

        skills = analysis_data.get('skills', [])
        for idx, skill in enumerate(skills, start=2):
            ws_skills[f'A{idx}'] = skill

        ws_skills.column_dimensions['A'].width = 30

        # Sheet 3: Matched Keywords (if analysis exists)
        analysis = analysis_data.get('analysis', {})
        if analysis:
            ws_matched = wb.create_sheet('Matched Keywords')
            ws_matched['A1'] = 'Keyword'
            ws_matched['A1'].fill = header_fill
            ws_matched['A1'].font = header_font

            matched = analysis.get('matched_skills', [])
            for idx, keyword in enumerate(matched, start=2):
                ws_matched[f'A{idx}'] = keyword

            ws_matched.column_dimensions['A'].width = 30

            # Sheet 4: Missing Keywords
            ws_missing = wb.create_sheet('Missing Keywords')
            ws_missing['A1'] = 'Keyword'
            ws_missing['A1'].fill = header_fill
            ws_missing['A1'].font = header_font

            missing = analysis.get('missing_skills', [])
            for idx, keyword in enumerate(missing, start=2):
                ws_missing[f'A{idx}'] = keyword

            ws_missing.column_dimensions['A'].width = 30

            # Sheet 5: Suggestions
            ws_suggestions = wb.create_sheet('Suggestions')
            ws_suggestions['A1'] = 'Priority'
            ws_suggestions['B1'] = 'Suggestion'
            ws_suggestions['A1'].fill = header_fill
            ws_suggestions['A1'].font = header_font
            ws_suggestions['B1'].fill = header_fill
            ws_suggestions['B1'].font = header_font

            suggestions = analysis.get('suggestions', [])
            for idx, suggestion in enumerate(suggestions, start=2):
                ws_suggestions[f'A{idx}'] = idx - 1
                ws_suggestions[f'B{idx}'] = suggestion

            ws_suggestions.column_dimensions['A'].width = 10
            ws_suggestions.column_dimensions['B'].width = 60

        # Save to buffer
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer
